"""Builds the sample files the browser tests upload: one philosophy lecture as PowerPoint,
PDF and Word, plus a photo of notes. Run: python3 tests/fixtures/make.py"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Pt
from docx import Document
from reportlab.lib.pagesizes import landscape, letter
from reportlab.pdfgen import canvas
from PIL import Image, ImageDraw, ImageFont

HERE = Path(__file__).parent
SLIDES = [
    ("Ancient Philosophy", ["Lecture 3", "From the Presocratics to Aristotle"]),
    ("The Presocratics", ["Thales: everything is water", "Heraclitus: everything flows", "Parmenides: change is an illusion"]),
    ("Socrates and the Elenchus", ["The Socratic method of questioning", "Know thyself", "The unexamined life is not worth living"]),
    ("Plato: The Theory of Forms", ["Forms are perfect abstract realities", "The allegory of the cave", "Justice in the Republic"]),
    ("Plato vs Aristotle on the Forms", ["Aristotle rejected separate Forms", "Form exists in matter", "Hylomorphism"]),
    ("Aristotle: Virtue Ethics", ["Eudaimonia is the highest good", "Virtue is a mean between extremes", "Habituation builds character"]),
    ("Hellenistic Schools", ["Stoicism", "Epicureanism", "Skepticism"]),
]

# PowerPoint
prs = Presentation()
for title, bullets in SLIDES:
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = title
    tf = s.placeholders[1].text_frame
    tf.text = bullets[0]
    for b in bullets[1:]:
        tf.add_paragraph().text = b
prs.save(HERE / "lecture.pptx")

# PDF: one slide per page, big title
c = canvas.Canvas(str(HERE / "lecture.pdf"), pagesize=landscape(letter))
W, H = landscape(letter)
for i, (title, bullets) in enumerate(SLIDES, 1):
    c.setFont("Helvetica-Bold", 34); c.drawString(60, H - 90, title)
    c.setFont("Helvetica", 20)
    for j, b in enumerate(bullets):
        c.drawString(80, H - 160 - j * 36, "• " + b)
    c.setFont("Helvetica", 12); c.drawString(W - 80, 30, f"{i} / {len(SLIDES)}")
    c.showPage()
c.save()

# Word: headings split it into sections
d = Document()
for title, bullets in SLIDES:
    d.add_heading(title, level=1)
    for b in bullets:
        d.add_paragraph(b)
d.save(HERE / "lecture.docx")

# A "photo" of notes for text recognition
img = Image.new("RGB", (1200, 700), (250, 248, 240))
dr = ImageDraw.Draw(img)
try:
    big = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", 56)
    small = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 36)
except OSError:
    big = small = ImageFont.load_default()
dr.text((60, 60), "Aristotle: Virtue Ethics", fill=(20, 20, 30), font=big)
for j, line in enumerate(["Eudaimonia is the highest good", "Virtue is a mean between extremes", "Habituation builds character"]):
    dr.text((80, 200 + j * 80), "- " + line, fill=(40, 40, 50), font=small)
img.save(HERE / "notes-photo.png")
print("fixtures written to", HERE)
